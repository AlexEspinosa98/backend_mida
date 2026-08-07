"""The single most important test in this package.

Validates apps.who_standards.lms.compute_zscore against WHO's OWN
pre-computed SD-column values, sampled directly from the raw downloaded
.xlsx files (apps/who_standards/data/raw/) -- NOT from our derived CSVs,
and NOT from any invented/memorized numbers. This proves both:
  1. the LMS formula implementation is correct, and
  2. nothing got mangled while loading the source files.

For a handful of rows (first/middle/last) in every one of the 18 raw WHO
source files, we assert that compute_zscore(SD{k}[neg], L, M, S) is very
close to +-k, for k in {1, 2, 3}.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from apps.who_standards.lms import compute_zscore

RAW_DIR = Path(__file__).parent.parent / "data" / "raw"

# WHO's published SD{k}[neg] columns are ROUNDED to 1-2 decimals (kg or cm),
# and the LMS curve's slope is steep enough near the +-3SD tails -- and for
# indicators with a small M (e.g. newborn weight-for-age, M~3.3kg) -- that
# this rounding alone produces z-score deviations well above a naive 1e-2.
#
# Empirically measured (see apps/who_standards/tests/ conversation history /
# PROVENANCE.md investigation): sampling first/middle/last rows across all
# 18 raw WHO files, the worst-case |z - expected_z| was ~0.25 (at SD2neg for
# a low-weight newborn wfl row), with a mean of ~0.02 across 324 sample
# points. A real formula bug (wrong sign, swapped L/S, wrong exponent)
# produces errors an order of magnitude larger and non-randomly distributed,
# not isolated to low-M extreme-SD points -- so:
#   - PER_POINT_TOLERANCE is loose enough to absorb WHO's own rounding.
#   - MEAN_TOLERANCE (aggregate, per file) is tight and catches systematic
#     bias that a real bug would introduce.
PER_POINT_TOLERANCE = 0.3
MEAN_TOLERANCE = 0.08


def _read_raw_rows(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() for c in rows[0]]
    data = [dict(zip(header, r)) for r in rows[1:]]
    wb.close()
    return data


def _sample_rows(rows: list[dict]) -> list[dict]:
    """first, middle, last -- deduplicated in case the table is tiny."""
    idxs = sorted({0, len(rows) // 2, len(rows) - 1})
    return [rows[i] for i in idxs]


ALL_RAW_FILES = sorted(RAW_DIR.glob("*.xlsx"))


@pytest.mark.parametrize("raw_path", ALL_RAW_FILES, ids=lambda p: p.name)
def test_sd_columns_match_lms_formula(raw_path: Path):
    rows = _read_raw_rows(raw_path)
    assert rows, f"{raw_path} appears empty"

    sd_to_z = {
        "SD3neg": -3.0,
        "SD2neg": -2.0,
        "SD1neg": -1.0,
        "SD1": 1.0,
        "SD2": 2.0,
        "SD3": 3.0,
    }

    abs_diffs = []
    for row in _sample_rows(rows):
        L, M, S = float(row["L"]), float(row["M"]), float(row["S"])
        for sd_col, expected_z in sd_to_z.items():
            x = row[sd_col]
            assert x is not None, f"{raw_path.name}: missing {sd_col} value"
            z = compute_zscore(float(x), L, M, S)
            abs_diffs.append(abs(z - expected_z))
            assert z == pytest.approx(expected_z, abs=PER_POINT_TOLERANCE), (
                f"{raw_path.name}: compute_zscore({x}, L={L}, M={M}, S={S}) "
                f"= {z}, expected ~{expected_z} (column {sd_col})"
            )

    mean_diff = sum(abs_diffs) / len(abs_diffs)
    assert mean_diff < MEAN_TOLERANCE, (
        f"{raw_path.name}: mean |z - expected_z| across sampled points = "
        f"{mean_diff:.4f}, expected < {MEAN_TOLERANCE} -- this suggests a "
        f"systematic bias (possible formula bug), not just WHO's rounding."
    )
